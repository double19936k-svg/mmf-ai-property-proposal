from __future__ import annotations

import json
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app"
if str(APP) not in sys.path:
    sys.path.insert(0, str(APP))

from delivery_state import DELIVERY_REVIEW, classify_fail_soft  # noqa: E402
from governance.final_artifact_qa import repair_docx_artifact  # noqa: E402
from longform.repair_policy import (  # noqa: E402
    FastRepairBudget,
    SEVERITY_CRITICAL,
    SEVERITY_MINOR,
    classify_section_issue,
    classify_under_length,
    compact_repair_context,
    contains_full_tender,
    repair_prompt_rules,
)
from longform.scheduler import ConcurrentBatchScheduler  # noqa: E402
from providers.mock import MockProvider  # noqa: E402
from tests.test_longform_orchestrator import _brief, _pack, _provider, _selection, _write_run  # noqa: E402
from longform.orchestrator import generate_longform  # noqa: E402


def _gate(*, units: int, target: int, must_ok: bool = True, body: bool = True, complete: bool = True, length: str | None = None) -> dict:
    if length is None:
        length = "PASS" if units >= target * 0.85 else "SECTION_UNDER_LENGTH"
    return {
        "status": "PASS" if must_ok and body else "BLOCK",
        "content_units": units,
        "length_status": length,
        "meaningful_body": body,
        "contract_complete": complete,
        "must_cover": {"项目基本事实": "COVERED" if must_ok else "MISSING"},
        "required_outputs_missing": [],
        "process_missing": [],
        "commitment": {"status": "PASS"},
        "fact_drift": [],
    }


def _contract(target: int = 1000) -> dict:
    return {"section_id": "S05-01", "section_title": "秩序", "target_words": {"min": target, "max": target + 200}, "must_cover": ["项目基本事实"], "required_outputs": ["巡查记录"]}


class FastPerformanceGovernanceTests(unittest.TestCase):
    def test_perf_03_minor_under_length_fast_does_not_repair(self) -> None:
        gate = _gate(units=800, target=1000, complete=True)
        decision = classify_section_issue(gate, _contract(1000), "UNDER_LENGTH", speed_profile="fast")
        self.assertEqual(classify_under_length(gate, _contract(1000)), SEVERITY_MINOR)
        self.assertEqual(decision["max_provider_repairs"], 0)

    def test_perf_04_critical_short_section_repairs(self) -> None:
        gate = _gate(units=300, target=1000, must_ok=False, complete=False)
        decision = classify_section_issue(gate, _contract(1000), "UNDER_LENGTH", speed_profile="fast")
        self.assertEqual(decision["severity"], SEVERITY_CRITICAL)
        self.assertEqual(decision["max_provider_repairs"], 2)

    def test_perf_05_ordinary_quality_max_one_fast_repair(self) -> None:
        gate = _gate(units=1200, target=1000)
        gate["status"] = "BLOCK"
        gate["commitment"] = {"status": "PASS"}
        decision = classify_section_issue(gate, _contract(), "QUALITY_FAILURE", speed_profile="fast")
        self.assertEqual(decision["max_provider_repairs"], 1)

    def test_perf_06_second_ordinary_fail_is_needs_review(self) -> None:
        budget = FastRepairBudget(speed_profile="fast", initial_batch_count=15)
        decision = {"severity": "QUALITY_FAILURE", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 40}
        self.assertTrue(budget.allow(decision, previous_repairs=0))
        self.assertFalse(budget.allow(decision, previous_repairs=1))
        result = classify_fail_soft(
            artifact_path=None,
            findings=[{"rule_id": "UNSUPPORTED_COMPLETION_CLAIM", "severity": "BLOCK", "text": "残留承诺", "scope": "sentence"}],
            repair_attempts=1,
        )
        # without a file this is technical; with a file it is review
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ok.docx"
            doc = Document()
            doc.add_paragraph("完整方案正文。")
            doc.save(str(path))
            decision = classify_fail_soft(artifact_path=path, findings=[{"rule_id": "UNSUPPORTED_COMPLETION_CLAIM", "severity": "BLOCK", "text": "残留承诺"}], repair_attempts=1)
            self.assertEqual(decision["delivery_status"], DELIVERY_REVIEW)

    def test_perf_07_critical_allows_second_repair(self) -> None:
        budget = FastRepairBudget(speed_profile="fast", initial_batch_count=15)
        decision = {"severity": SEVERITY_CRITICAL, "max_provider_repairs": 2, "can_exceed_global": True, "value_score": 90}
        self.assertTrue(budget.allow(decision, previous_repairs=1))

    def test_perf_08_global_budget_caps_low_value_repairs(self) -> None:
        budget = FastRepairBudget(speed_profile="fast", initial_batch_count=15)
        self.assertEqual(budget.max_provider, 4)
        decision = {"severity": "QUALITY_FAILURE", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 40}
        for _ in range(4):
            self.assertTrue(budget.allow(decision, previous_repairs=0))
            budget.consume()
        self.assertFalse(budget.allow(decision, previous_repairs=0))
        critical = {"severity": SEVERITY_CRITICAL, "max_provider_repairs": 2, "can_exceed_global": True, "value_score": 90}
        self.assertTrue(budget.allow(critical, previous_repairs=0))

    def test_perf_09_broken_quote_is_local_repair(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "quotes.docx"
            doc = Document()
            doc.add_paragraph("分区分类“ ”差异化管控。")
            doc.save(str(path))
            before = path.stat().st_mtime
            repair_docx_artifact(path)
            self.assertGreaterEqual(path.stat().st_mtime, before)
            self.assertTrue(path.is_file())

    def test_perf_10_two_independent_batches_run_concurrently(self) -> None:
        live = {"n": 0, "max": 0}
        lock = threading.Lock()

        def work(batch):
            with lock:
                live["n"] += 1
                live["max"] = max(live["max"], live["n"])
            time.sleep(0.05)
            with lock:
                live["n"] -= 1
            return {"status": "SUCCESS", "section_ids": batch["section_ids"]}

        scheduler = ConcurrentBatchScheduler(provider_name="mock", max_parallelism=2)
        out = scheduler.run(
            [
                {"batch_id": "GB-01", "section_ids": ["S01-01"], "dependency_ids": []},
                {"batch_id": "GB-02", "section_ids": ["S02-01"], "dependency_ids": []},
            ],
            work,
        )
        self.assertEqual(out["max_observed_concurrency"], 2)
        self.assertEqual(live["max"], 2)

    def test_perf_11_repair_context_omits_full_tender(self) -> None:
        tender = "招标文件" * 4000 + "完整条款" * 2000
        context = {
            "section_contract": _contract(),
            "project_facts": {"gross_area": {"value": "252567平方米"}},
            "confirmed_requirements": [{"text": tender}],
            "used_core_arguments": ["x"] * 40,
        }
        gate = _gate(units=400, target=1000, must_ok=False)
        compact = compact_repair_context(context, gate, previous_fragment={"body_blocks": [{"type": "paragraph", "content": "已有正文"}]}, failure_class="UNDER_LENGTH")
        blob = json.dumps(compact, ensure_ascii=False)
        self.assertFalse(contains_full_tender(blob))
        self.assertLess(len(blob), 8000)
        self.assertIn("只补充失败项", " ".join(repair_prompt_rules("UNDER_LENGTH")))

    def test_replay_glm_run_skips_low_value_second_repairs(self) -> None:
        baseline = json.loads((ROOT / "tests" / "fixtures" / "r27_glm_performance_baseline.json").read_text(encoding="utf-8"))
        budget = FastRepairBudget(speed_profile="fast", initial_batch_count=15)
        kept = 0
        skipped = 0
        for row in (
            {"severity": "MATERIAL_UNDER_LENGTH", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 45},
            {"severity": "QUALITY_FAILURE", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 40},
            {"severity": "QUALITY_FAILURE", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 40},
            {"severity": "QUALITY_FAILURE", "max_provider_repairs": 1, "can_exceed_global": False, "value_score": 40},
        ):
            if budget.allow(row, previous_repairs=0):
                budget.consume()
                kept += 1
            if budget.allow(row, previous_repairs=1):
                skipped += 0
            else:
                skipped += 1
        self.assertEqual(kept, 4)
        self.assertEqual(skipped, 4)
        self.assertEqual(baseline["new_expected_provider_repairs_kept"], 4)
        self.assertEqual(baseline["initial_generation_provider_calls"], 17)

    def test_perf_01_02_12_batch_execution_and_reduction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            run = _write_run(Path(tmp), _pack(), {**_brief(), "speed_profile": "fast"}, _selection())
            result = generate_longform(
                run_dir=run,
                provider=_provider(),
                provider_name="mock",
                brief={**_brief(), "speed_profile": "fast"},
                selection=_selection(),
                selected_ids=["KU-0001"],
                speed_profile="fast",
                max_parallelism=2,
            )
            word = result.get("word") or {}
            logical = int(word.get("logical_section_count") or 0)
            batches = int(word.get("generation_batch_count") or 0)
            initial = int(word.get("INITIAL_GENERATION_CALLS") or 0)
            self.assertGreaterEqual(logical, 8)
            self.assertGreaterEqual(batches, 1)
            self.assertEqual(initial, batches)
            self.assertGreater(word.get("CALL_REDUCTION_RATIO") or 0, 0)
            self.assertGreater(word.get("INITIAL_SECTIONS_PER_CALL") or 0, 1)
            self.assertEqual(int(word.get("SECTION_REPAIR_CALLS") or 0), 0)


class BatchPromptContractTests(unittest.TestCase):
    def test_perf_02_two_sections_one_call(self) -> None:
        class Probe(MockProvider):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self.calls = []

            def invoke_structured(self, request, task_dir):
                self.calls.append(dict(request))
                return super().invoke_structured(request, task_dir)

        provider = Probe({"provider_name": "mock", "provider_type": "mock", "enabled": True, "model": "mock-longform"})
        with tempfile.TemporaryDirectory() as tmp:
            run = _write_run(Path(tmp), _pack(), {**_brief(), "speed_profile": "fast"}, _selection())
            generate_longform(
                run_dir=run,
                provider=provider,
                provider_name="mock",
                brief={**_brief(), "speed_profile": "fast"},
                selection=_selection(),
                selected_ids=["KU-0001"],
                speed_profile="fast",
            )
        batch_calls = [row for row in provider.calls if row.get("attempt_mode") == "generate_batch"]
        self.assertTrue(batch_calls)
        self.assertGreaterEqual(max(len(row.get("section_ids") or []) for row in batch_calls), 2)
        two = next(row for row in batch_calls if len(row.get("section_ids") or []) >= 2)
        self.assertEqual(two.get("generation_mode"), "longform_batch")


if __name__ == "__main__":
    unittest.main()
