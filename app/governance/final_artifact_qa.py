from __future__ import annotations

import hashlib
import os
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .artifact_qa import evaluate_artifact
from .artifact_qa import normalize_quotes
from .commitment_provenance import evaluate_commitments
from .customer_hygiene import hygiene_findings, rewrite_customer_hygiene


def _heading_level(style_name: str) -> int | None:
    name = str(style_name or "").strip()
    match = re.search(r"(?:Heading|标题)\s*([1-9])", name, re.I)
    return int(match.group(1)) if match else None


def _looks_like_heading(text: str, style_name: str) -> bool:
    compact = re.sub(r"\s+", "", str(text or ""))
    if _heading_level(style_name) is not None:
        return True
    if str(style_name or "").lower().startswith("list"):
        return False
    return bool(4 <= len(compact) <= 30 and not re.search(r"[，。；：！？,.!?;:]", compact))


def _docx_blocks(path: Path) -> list[dict[str, Any]]:
    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            text = str(paragraph.text or "").strip()
            style = str(paragraph.style.name if paragraph.style is not None else "")
            if text:
                blocks.append({
                    "path": f"paragraph[{paragraph_index}]", "paragraph_index": paragraph_index,
                    "text": text, "kind": "paragraph", "heading_level": _heading_level(style),
                    "heading_candidate": _looks_like_heading(text, style), "style": style,
                })
            paragraph_index += 1
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            text = "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows).strip()
            if text:
                blocks.append({"path": f"table[{table_index}]", "text": text, "kind": "table", "heading_level": None, "style": "Table"})
            table_index += 1
    return blocks


def _empty_headings(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    for index, row in enumerate(blocks):
        level = row.get("heading_level")
        if not row.get("heading_candidate") or row["text"].strip() in {"目录", "目 录"}:
            continue
        following_rows = blocks[index + 1:]
        if following_rows and following_rows[0].get("heading_candidate"):
            findings.append({
                "rule_id": "EMPTY_HEADING_BLOCK",
                "path": row["path"],
                "severity": "BLOCK",
                "issue": f"heading is immediately followed by another heading: {row['text']}",
                "text": row["text"],
                "repair_action": "REMOVE_EMPTY_HEADING",
            })
            continue
        if level is None:
            continue
        meaningful = False
        for following in blocks[index + 1:]:
            next_level = following.get("heading_level")
            if next_level is not None and next_level <= level:
                break
            if following.get("kind") == "table" or (next_level is None and len(re.sub(r"\s+", "", following["text"])) >= 2):
                meaningful = True
                break
        if not meaningful:
            findings.append({
                "rule_id": "EMPTY_HEADING_BLOCK",
                "path": row["path"],
                "severity": "BLOCK",
                "issue": f"heading has no meaningful body: {row['text']}",
            })
    return findings


INTERNAL_SECTION_ID = re.compile(r"(?<![A-Za-z0-9])S\d{2}-\d{2}(?![A-Za-z0-9])")
EMPTY_QUOTE = re.compile(r"[“‘][\s\u3000]*[”’]")
ADJACENT_QUOTE = re.compile(r"[”’][\s\u3000]*[“‘]")
VALID_DOUBLE_QUOTE = re.compile(r"“[^”\n]{1,120}”")
UNSUPPORTED_ARTIFACT_CLAIMS = (
    re.compile(r"(?:所有|全部)[^。；\n]{0,32}(?:履历|资历证明)[^。；\n]{0,16}(?:将|已)[^。；\n]{0,8}(?:单独成册|另行成册|整理成册)"),
    re.compile(r"(?:简历|履历|资历评价)[^。；\n]{0,16}(?:详见|见)[^。；\n]{0,8}(?:附表|附件)"),
    re.compile(r"所有拟派人员[^。；\n]{0,12}(?:已经|均已)[^。；\n]{0,24}"),
)
RECOVERABLE_RULES = {
    "QUOTE_BALANCE_QA", "BROKEN_QUOTE_SERIALIZATION", "EMPTY_HEADING_BLOCK",
    "UNSUPPORTED_COMPLETION_CLAIM", "INTERNAL_SECTION_ID",
    "CONDITIONAL_REWRITE_GRAMMAR", "BROKEN_CROSS_REFERENCE", "BROKEN_FLOW_SERIALIZATION",
    "UNSUPPORTED_SYSTEM_CAPABILITY", "UNSUPPORTED_NUMERIC_THRESHOLD", "INTERNAL_GOVERNANCE_TONE",
    "CONDITIONAL_POSTPROCESS_POLLUTION", "CLIENT_ARTIFACT_INTERNAL_META",
}


def quote_serialization_status(payload: Any) -> dict[str, Any]:
    """Return an auditable quote status for provider/intermediate JSON text."""
    texts: list[str] = []

    def visit(value: Any) -> None:
        if isinstance(value, str):
            texts.append(value)
        elif isinstance(value, dict):
            for child in value.values():
                visit(child)
        elif isinstance(value, (list, tuple)):
            for child in value:
                visit(child)

    visit(payload)
    exact = sum(len(EMPTY_QUOTE.findall(text)) + len(ADJACENT_QUOTE.findall(text)) for text in texts)
    unbalanced = sum(int(text.count("“") != text.count("”") or text.count("‘") != text.count("’")) for text in texts)
    valid_pairs = sum(len(VALID_DOUBLE_QUOTE.findall(text)) for text in texts)
    status = "FAIL" if exact or unbalanced else ("RISK" if valid_pairs >= 20 else "PASS")
    return {
        "status": status,
        "empty_or_adjacent_count": exact,
        "unbalanced_text_count": unbalanced,
        "curly_quote_pair_count": valid_pairs,
    }


def _quote_findings(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    valid_pairs = 0
    quote_paths: list[str] = []
    for row in blocks:
        text = row["text"]
        exact = len(EMPTY_QUOTE.findall(text)) + len(ADJACENT_QUOTE.findall(text))
        unbalanced = int(text.count("“") != text.count("”") or text.count("‘") != text.count("’"))
        if exact or unbalanced:
            findings.append({
                "rule_id": "BROKEN_QUOTE_SERIALIZATION", "path": row["path"], "severity": "BLOCK",
                "issue": "empty, adjacent, or unbalanced Chinese quote marks in final DOCX",
                "occurrence_count": exact + unbalanced, "repair_action": "NORMALIZE_QUOTES",
            })
        pairs = len(VALID_DOUBLE_QUOTE.findall(text))
        if pairs:
            valid_pairs += pairs
            quote_paths.append(row["path"])
    # A high concentration of curly quotes in the legacy renderer is a known WPS
    # compatibility failure: the enclosed text can render outside the marks and
    # leave visually empty quotes even though plain-text extraction looks balanced.
    if valid_pairs >= 20:
        findings.append({
            "rule_id": "BROKEN_QUOTE_SERIALIZATION", "path": quote_paths[0] if quote_paths else "artifact",
            "severity": "BLOCK", "issue": "WPS curly-quote serialization compatibility risk",
            "occurrence_count": valid_pairs, "repair_action": "NORMALIZE_QUOTES_FOR_WPS",
        })
    return findings


def _deterministic_artifact_findings(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    findings = _quote_findings(blocks)
    for row in blocks:
        text = row["text"]
        ids = INTERNAL_SECTION_ID.findall(text)
        if ids:
            findings.append({
                "rule_id": "INTERNAL_SECTION_ID", "path": row["path"], "severity": "BLOCK",
                "issue": "internal section id leaked into customer artifact", "text": "、".join(ids),
                "occurrence_count": len(ids), "repair_action": "REMOVE_INTERNAL_SECTION_ID",
            })
        for pattern in UNSUPPORTED_ARTIFACT_CLAIMS:
            matches = list(pattern.finditer(text))
            if matches:
                findings.append({
                    "rule_id": "UNSUPPORTED_COMPLETION_CLAIM", "path": row["path"], "severity": "BLOCK",
                    "issue": "completion or attachment claim lacks an evidence registry entry",
                    "text": matches[0].group(0), "occurrence_count": len(matches),
                    "repair_action": "REWRITE_AS_REQUIRED_FUTURE_EVIDENCE",
                })
        findings.extend(hygiene_findings(text, row["path"], severity="BLOCK"))
    return findings


def artifact_findings_are_locally_recoverable(report: dict[str, Any]) -> bool:
    blocked = [row for row in (report.get("findings") or []) if row.get("severity") == "BLOCK"]
    return bool(blocked) and all(str(row.get("rule_id")) in RECOVERABLE_RULES for row in blocked)


def _replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def repair_docx_artifact(path: Path) -> dict[str, Any]:
    """Apply deterministic, provider-independent repairs to a newly rendered DOCX."""
    target = Path(path)
    before = _docx_blocks(target)
    empty_indices = sorted({int(row["path"][10:-1]) for row in _empty_headings(before)}, reverse=True)
    quote_pair_count = sum(len(VALID_DOUBLE_QUOTE.findall(row["text"])) for row in before)
    doc = Document(str(target))
    counts = {"quote": 0, "empty_heading": 0, "unsupported_completion": 0, "internal_section_id": 0}
    for paragraph in doc.paragraphs:
        original = paragraph.text or ""
        fixed = normalize_quotes(original)
        if quote_pair_count >= 20:
            fixed = fixed.replace("“", "「").replace("”", "」").replace("‘", "『").replace("’", "』")
        counts["quote"] += int(fixed != original and any(ch in original for ch in "“”‘’"))
        fixed, count = re.subn(r"S\d{2}-\d{2}\s*[（(]([^）)]{1,30})[）)]", r"\1", fixed)
        counts["internal_section_id"] += count
        before_hygiene = fixed
        fixed = rewrite_customer_hygiene(fixed)
        counts["internal_section_id"] += int("S0" in before_hygiene and "S0" not in fixed)
        replacements = (
            (re.compile(r"(?:注：)?所有岗位人员履历及资历证明将单独成册[^。；\n]*[。；]?"), "拟派团队资历及相关证明材料应按招标要求另行提供。"),
            (re.compile(r"其简历及资历评价详见附表"), "拟派人员资历及相关证明材料应按招标要求另行提供"),
            (re.compile(r"所有拟派人员(?:已经|均已)[^。；\n]*"), "拟派人员安排及资历证明应以实际提交并经确认的材料为准"),
        )
        for pattern, replacement in replacements:
            fixed, changed = pattern.subn(replacement, fixed)
            counts["unsupported_completion"] += changed
        if fixed != original:
            _replace_paragraph_text(paragraph, fixed)
    for index in empty_indices:
        if 0 <= index < len(doc.paragraphs):
            element = doc.paragraphs[index]._element
            element.getparent().remove(element)
            counts["empty_heading"] += 1
    tmp = target.with_name(f"{target.stem}.artifact-qa-repair.tmp{target.suffix}")
    doc.save(str(tmp))
    os.replace(tmp, target)
    return {"status": "repaired", "artifact_path": str(target), "repairs": counts}


def evaluate_docx_artifact(
    path: Path,
    *,
    canonical_facts: dict[str, Any] | None = None,
    brief: dict[str, Any] | None = None,
    contracts: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    target = Path(path)
    if not target.is_file() or target.suffix.lower() != ".docx":
        return {
            "schema_version": "final-docx-artifact-qa-v0.1",
            "status": "BLOCK",
            "actual_docx_inspected": False,
            "findings": [{"rule_id": "FINAL_ARTIFACT_MISSING", "path": str(target), "severity": "BLOCK", "issue": "final DOCX is missing"}],
            "summary": {"block_count": 1, "warning_count": 0},
        }
    blocks = _docx_blocks(target)
    texts = [row["text"] for row in blocks]
    structured = evaluate_artifact({"artifact": {"paragraphs": texts}}, canonical_facts=canonical_facts)
    findings = list(structured.get("findings") or [])
    for item in findings:
        if item.get("rule_id") == "CLIENT_ARTIFACT_INTERNAL_META":
            item["severity"] = "BLOCK"
    findings.extend(_empty_headings(blocks))
    findings.extend(_deterministic_artifact_findings(blocks))
    commitment = evaluate_commitments(brief or canonical_facts or {}, contracts or [], {"artifact": {"paragraphs": texts}})
    findings.extend({
        "rule_id": str(row.get("claim_type") or "UNSUPPORTED_COMMITMENT"),
        "path": row.get("path") or "artifact",
        "severity": "BLOCK",
        "issue": row.get("reason") or "unsupported commitment",
        "text": row.get("text"),
        "repair_action": row.get("repair_action"),
    } for row in (commitment.get("violations") or []))
    unique: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for item in findings:
        key = (str(item.get("rule_id")), str(item.get("path")), str(item.get("issue")))
        if key not in seen:
            seen.add(key)
            unique.append(item)
    blocks_count = sum(item.get("severity") == "BLOCK" for item in unique)
    warnings = sum(item.get("severity") != "BLOCK" for item in unique)
    return {
        "schema_version": "final-docx-artifact-qa-v0.2",
        "status": "BLOCK" if blocks_count else ("WARNING" if warnings else "PASS"),
        "actual_docx_inspected": True,
        "artifact_path": str(target),
        "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "paragraph_count": sum(row["kind"] == "paragraph" for row in blocks),
        "table_count": sum(row["kind"] == "table" for row in blocks),
        "findings": unique,
        "summary": {"block_count": blocks_count, "warning_count": warnings},
        "quote_status": "FAIL" if any(row.get("rule_id") == "BROKEN_QUOTE_SERIALIZATION" for row in unique) else "PASS",
        "broken_quote_count": sum(int(row.get("occurrence_count") or 0) for row in unique if row.get("rule_id") == "BROKEN_QUOTE_SERIALIZATION"),
    }
